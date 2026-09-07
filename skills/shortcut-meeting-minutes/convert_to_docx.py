#!/usr/bin/env python3
import sys
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

# 參數
md_file = sys.argv[1] if len(sys.argv) > 1 else '/tmp/meeting_minutes.md'
docx_file = sys.argv[2] if len(sys.argv) > 2 else md_file.replace('.md', '.docx')

# 讀取 Markdown
with open(md_file, 'r', encoding='utf-8') as f:
    lines = f.readlines()

doc = Document()

# 標題
title_text = '會議記錄: IT Team Review'
for line in lines:
    if line.startswith('會議記錄:'):
        title_text = line.strip()
        break

title = doc.add_heading(title_text, level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in title.runs:
    run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
    run.font.size = Pt(24)

doc.add_paragraph()

# 處理內容
skip_header = True
for line in lines:
    line = line.rstrip('\n')

    if skip_header:
        if line.startswith('## 大綱'):
            skip_header = False
            doc.add_heading('大綱', level=1)
        continue

    if not line.strip() or line.startswith('===') or line.startswith('---'):
        continue

    if line.startswith('### '):
        heading = doc.add_heading(line[4:], level=2)
        for run in heading.runs:
            run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
    elif line.startswith('## '):
        doc.add_heading(line[3:], level=1)
    elif line.startswith('- '):
        item_text = line[2:]
        p = doc.add_paragraph(item_text, style='List Bullet')
        p.paragraph_format.left_indent = Inches(0.25)
        for run in p.runs:
            run.font.size = Pt(10)
    elif line.strip():
        p = doc.add_paragraph(line)
        for run in p.runs:
            run.font.size = Pt(11)

doc.save(docx_file)

file_size = os.path.getsize(docx_file)
print(f"{docx_file}|{file_size}")
